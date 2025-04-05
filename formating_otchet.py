import os
import json
import pandas as pd
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

FOLDERS = [
    "1_курс", "2_курс", "3_курс",
    "4_курс", "5_курс",
    "1_курс_мага", "2_курс_мага"
]


def add_empty_rows(df, num_rows=6):
    wb = Workbook()
    ws = wb.active
    for _ in range(num_rows):
        ws.append([])
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)
    return wb


def add_names_columns(df):
    if "М1" not in df.columns:
        df["М1"] = ""
    if "М2" not in df.columns:
        df["М2"] = ""
    if "Примечание" not in df.columns:
        df["Примечание"] = ""
    return df


def combine_columns(df):
    cols = {"Фамилия", "Имя", "Отчество"}
    if cols.issubset(df.columns):
        for c in ["Фамилия", "Имя", "Отчество"]:
            df[c] = df[c].astype(str).str.strip()
        df["ФИО"] = (df["Фамилия"] + " " + df["Имя"] + " " + df["Отчество"]).str.strip()
        df.drop(["Фамилия", "Имя", "Отчество"], axis=1, inplace=True)
    return df


def process_dataframe(df):
    if df.columns[0].startswith("Unnamed") or df.columns[0] == df.index.name:
        df.drop(df.columns[0], axis=1, inplace=True)
    df = combine_columns(df)
    df = add_names_columns(df)
    return df


def shorten_subject(subject):
    words = subject.strip().split()
    letters = []
    for i, w in enumerate(words):
        if w:
            if i == 0:
                letters.append(w[0].upper())
            else:
                letters.append(w[0].lower())
    return "".join(letters)


def set_worksheet_formats(ws):
    ws.merge_cells("A1:H1")
    row_h = 7
    max_col = ws.max_column
    for r in range(row_h, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        if all((v is None or str(v).strip() == "") for v in vals):
            break
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
    for col in range(1, ws.max_column + 1):
        length = 0
        col_letter = get_column_letter(col)
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            v = cell.value
            if v is not None:
                l = len(str(v))
                if l > length:
                    length = l
        ws.column_dimensions[col_letter].width = length + 2


def combined_excel_files():
    path = "new_data"
    os.makedirs(path, exist_ok=True)
    output = Path(path) / "combined.xlsx"
    if output.exists():
        output.unlink()
    with open("subject_to_files.json", "r", encoding="utf-8") as f:
        subject_groups = json.load(f)
    with open("scripts/group_directions.json", "r", encoding="utf-8") as f:
        group_dirs = json.load(f)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for folder in FOLDERS:
            folder_path = Path("scripts/data") / folder
            if not folder_path.exists():
                continue
            excel_files = folder_path.glob("*.xlsx")
            for file_path in excel_files:
                df = pd.read_excel(file_path)
                df = process_dataframe(df)
                for subject, groups in subject_groups.items():
                    for group in groups:
                        short_subj = shorten_subject(subject)
                        sheet_name = f"{short_subj} {group}"
                        df.to_excel(writer, sheet_name=sheet_name, startrow=6, index=False)
                        ws = writer.sheets[sheet_name]
                        ws["A1"] = subject
                        ws["H3"] = group
                        direction_value = group_dirs.get(group, "")
                        ws["H4"] = direction_value
                        ws["B3"] = "Лектор"
                        ws["B4"] = "Семинар"
                        ws["B5"] = "Лабораторные"
    wb = load_workbook(output)
    for ws in wb.worksheets:
        set_worksheet_formats(ws)
    wb.save(output)
    print("Done", output)


def main():
    combined_excel_files()


if __name__ == "__main__":
    main()
