import os
import json
from openpyxl import load_workbook, Workbook
from copy import copy
from openpyxl.cell import Cell

script_dir = os.path.dirname(os.path.abspath(__file__))
base_path = os.path.join(script_dir, 'data')
subfolders = ["1_курс", "2_курс", "3_курс", "4_курс", "5_курс", '1_курс_мага', '2_курс_мага']
os.remove('group_directions.json')
group_keywords = ['МДС', 'ИДБ', 'ЭДБ', 'АДБ', 'МДБ', 'ИДМ', 'МДМ', 'ЭДМ', 'АДМ']

group_directions = {}

for subfolder in subfolders:
    folder_path = os.path.join(base_path, subfolder)
    if not os.path.exists(folder_path):
        print(f"Папка {folder_path} не найдена, пропускаем.")
        continue

    for root, _, files in os.walk(folder_path):
        for file in files:
            if file.endswith('.xlsx'):
                file_path = os.path.join(root, file)
                print(f"Обработка файла: {file_path}")
                try:
                    wb = load_workbook(file_path)
                    ws = wb.active
                    header = [cell.value if isinstance(cell, Cell) else cell for cell in ws[1]]

                    group_columns = [col for col in header if
                                     isinstance(col, str) and any(key in col for key in group_keywords)]

                    output_folder = os.path.dirname(file_path)
                    for group in group_columns:
                        group_index = header.index(group)
                        # Для json файла ключ-значение (группа: направление)
                        direction_index = group_index + 2

                        if direction_index < len(header):
                            direction = header[direction_index]
                        else:
                            direction = "Не указано"
                        group_directions[group] = direction

                        group_cols = [group_index, group_index + 1, group_index + 2]
                        group_data = []
                        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                            if any(cell.value for cell in [row[i] for i in group_cols]):
                                group_data.append(row)
                        if not group_data:
                            continue
                        new_wb = Workbook()
                        new_ws = new_wb.active
                        new_ws.title = group
                        new_ws.append(['№', 'Фамилия', 'Имя', 'Отчество'])
                        for idx, row in enumerate(group_data, start=1):
                            new_row = [idx] + [row[i].value for i in group_cols]
                            new_ws.append(new_row)
                            for source_idx, target_cell in zip(group_cols, new_ws[idx + 1][1:]):
                                source_cell = row[source_idx]
                                target_cell.font = copy(source_cell.font)
                                target_cell.alignment = copy(source_cell.alignment)
                                target_cell.border = copy(source_cell.border)
                                target_cell.fill = copy(source_cell.fill)
                        output_path = os.path.join(output_folder, f"{group}.xlsx")
                        new_wb.save(output_path)
                        # print(f"Данные для группы {group} сохранены в файл {output_path}")
                    os.remove(file_path)
                    # print(f"Файл {file_path} удалён после обработки.")
                except Exception as e:
                    print(f"Ошибка при обработке файла {file_path}: {e}")

group_dir_json = os.path.join(script_dir, 'group_directions.json')
with open(group_dir_json, 'w', encoding='utf-8') as file:
    json.dump(group_directions, file, ensure_ascii=False, indent=4)
