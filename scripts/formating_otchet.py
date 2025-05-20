import os
import json
import pandas as pd
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment
from openpyxl.styles import Border, Side

FOLDERS = [
    "1_курс", "2_курс", "3_курс",
    "4_курс", "5_курс",
    "1_курс_мага", "2_курс_мага"
]


def add_names_columns(df):
    for col in ("М1", "М2", "Примечание"):
        if col not in df.columns:
            df[col] = ""
    return df


def combine_columns(df):
    if "Фамилия" in df.columns and "Имя" in df.columns and "Отчество" in df.columns:
        def create_fio(row):
            last = str(row['Фамилия']).strip()
            first = str(row['Имя']).strip() if pd.notna(row['Имя']) else ""
            patronymic = str(row['Отчество']).strip() if pd.notna(row['Отчество']) else ""
            if first and patronymic:
                return f"{last} {first[0].upper()}. {patronymic[0].upper()}."
            elif first:
                return f"{last} {first[0].upper()}."
            elif patronymic:
                return f"{last} {patronymic[0].upper()}."
            else:
                return last

        df["ФИО"] = df.apply(create_fio, axis=1)
        df.drop(["Фамилия", "Имя", "Отчество"], axis=1, inplace=True)
    return df


def process_dataframe(df):
    first = df.columns[0]
    if first.startswith("Unnamed") or first == df.index.name:
        df.rename(columns={first: '№'}, inplace=True)
    if '№' not in df.columns:
        df.insert(0, '№', range(1, len(df) + 1))

    df = combine_columns(df)
    df = add_names_columns(df)

    df = df[['№', 'ФИО', 'М1', 'М2', 'Примечание']]

    df.insert(2, '', '')
    return df


def shorten_subject(subject):
    letters = [
        w[0].upper() if i == 0 else w[0].lower()
        for i, w in enumerate(subject.split()) if w
    ]
    return ''.join(letters)


def set_worksheet_formats(ws):
    ws.merge_cells("A1:H1")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    row_h = 7
    max_col = ws.max_column
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_h, column=col)
        cell.border = thin_border

    for r in range(row_h + 1, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        if all(v is None or str(v).strip() == "" for v in vals):
            break
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).border = thin_border
    for r in range(row_h, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        if all(v is None or str(v).strip() == "" for v in vals):
            break
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)

    ws.column_dimensions['A'].width = 5

    for col in range(2, max_col + 1):
        letter = get_column_letter(col)
        length = max(
            (
                len(str(ws.cell(row=r, column=col).value))
                for r in range(1, ws.max_row + 1)
                if ws.cell(row=r, column=col).value is not None
            ),
            default=0
        )
        ws.column_dimensions[letter].width = length + 2


def combined_excel_files():
    data_root = Path("data")
    output_dir = Path("../new_data")
    output_dir.mkdir(exist_ok=True)
    output = output_dir / "mod_vesn.xlsx"
    if output.exists(): output.unlink()

    with open("subject_to_files.json", encoding="utf-8") as f:
        subject_groups = json.load(f)
    with open("group_directions.json", encoding="utf-8") as f:
        group_dirs = json.load(f)
    with open("subject_to_files_extended.json", encoding="utf-8") as f:
        extended = json.load(f)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:

        for subject, groups in subject_groups.items():
            for group in groups:

                file_path = None
                for folder in FOLDERS:
                    candidate = data_root / folder / f"{group}.xlsx"
                    if candidate.exists():
                        file_path = candidate
                        break
                if not file_path:
                    continue

                wb_src = load_workbook(file_path)
                ws_src = wb_src.active
                hdr = next(ws_src.iter_rows(min_row=1, max_row=1))
                idx = {cell.value: cell.column for cell in hdr}
                fam_i, name_i = idx.get('Фамилия'), idx.get('Имя')
                bold_flags = []
                for row in ws_src.iter_rows(min_row=2, max_row=ws_src.max_row):
                    bold = False
                    if fam_i and row[fam_i - 1].font.bold: bold = True
                    if name_i and row[name_i - 1].font.bold: bold = True
                    bold_flags.append(bold)

                df = pd.read_excel(file_path)
                df = process_dataframe(df)

                sheet = f"{shorten_subject(subject)} {group}"
                df.to_excel(writer, sheet_name=sheet, startrow=6, startcol=0, index=False)
                ws = writer.sheets[sheet]

                ws["A1"] = subject
                ws["H3"] = group
                ws["H4"] = group_dirs.get(group, "")
                ws["B3"] = "Лектор"
                ws["B4"] = "Семинар"
                ws["B5"] = "Лабораторные"

                for e in extended.get(subject, []):
                    t, teacher, grp = e.get('type'), e.get('teacher'), e.get('group')
                    if t == 'Лекции':
                        ws['C3'] = teacher
                    elif grp.startswith(group):
                        if t == 'Практические': ws['C4'] = teacher
                        if t == 'Лабораторные': ws['C5'] = teacher

                start = 8
                for i, b in enumerate(bold_flags[:len(df)]):
                    if b:
                        ws.cell(row=start + i, column=2).font = Font(bold=True)

        if not writer.sheets:
            tmp = writer.book.create_sheet('TMP')
            writer.sheets['TMP'] = tmp

    wb = load_workbook(output)
    if 'TMP' in wb.sheetnames:
        wb.remove(wb['TMP'])
    for ws in wb.worksheets:
        set_worksheet_formats(ws)
    wb.save(output)
    print("Done", output)


def main():
    combined_excel_files()


if __name__ == '__main__':
    main()
